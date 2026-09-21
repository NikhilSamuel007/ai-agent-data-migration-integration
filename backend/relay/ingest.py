import csv
import io
from datetime import date, datetime
from zipfile import ZipFile
import openpyxl

def ingest(name, content):
    if len(content) > 5 * 1024 * 1024:
        raise ValueError('Each file must be at most 5 MB.')
    rows = []
    if name.lower().endswith('.csv'):
        reader = csv.reader(io.StringIO(content.decode('utf-8-sig'), newline=''), strict=True)
        for row in reader:
            if row:
                rows.append(row)
            if len(rows) > 1001:
                raise ValueError('Prototype limit: 1,000 records per file.')
    elif name.lower().endswith('.xlsx'):
        with ZipFile(io.BytesIO(content)) as archive:
            if sum(f.file_size for f in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError('Expanded Excel workbook exceeds 50 MB.')
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        try:
            if len(book.worksheets) != 1:
                raise ValueError('Use one worksheet per file so no sheet is silently ignored.')
            sheet = book.worksheets[0]
            if (sheet.max_column or 0) > 50 or (sheet.max_row or 0) > 1001:
                raise ValueError('Prototype limit: 50 columns and 1,000 records per file.')
            for cells in sheet.iter_rows():
                row = []
                for cell in cells:
                    if cell.data_type == 'f':
                        raise ValueError('Replace spreadsheet formulas with values before uploading.')
                    value = cell.value
                    if isinstance(value, datetime):
                        value = value.date().isoformat()
                    elif isinstance(value, date):
                        value = value.isoformat()
                    row.append('' if value is None else str(value))
                rows.append(row)
        finally:
            book.close()
    else:
        raise ValueError('Upload CSV or .xlsx files. Legacy .xls files must be saved as .xlsx.')
    if len(rows) < 2:
        raise ValueError(f'{name}: at least one header and one data row are required.')
    headers = [h.strip() for h in rows[0]]
    if any(not h for h in headers) or len(set(h.lower() for h in headers)) != len(headers):
        raise ValueError(f'{name}: headers must be non-empty and unique.')
    if len(headers) > 50:
        raise ValueError('Prototype limit: 50 columns per file.')
    if any(len(row) != len(headers) for row in rows[1:]):
        raise ValueError('Every row must have the same number of columns as the header.')
    return dict(name=name, headers=headers, rows=[dict(zip(headers, row)) for row in rows[1:]])
